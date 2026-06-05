#!/usr/bin/env python3
"""
Shared WooCommerce REST helpers.

This module is the WooCommerce-specific toolkit for reporting routines: an
authenticated, paginated REST client with retry/backoff, tolerant WC `meta_data`
parsing, shop-currency detection, ISO-week reporting windows, a Dropbox upload
and Excel styling. It lives in the shared `ClaudeWooCommerceCommons` package
(distribution `claude-woocommerce-commons`, import `wc_client`).

The generic, vendor-agnostic helpers — `env_required` / `env_opt` / `env_get`,
`parse_num`, `currency_symbol`, `build_remote_path`, `log` — live in the
dependency-light `claude-code-commons` package and are **re-exported here**, so
existing `from wc_client import env_required, parse_num, ...` keeps working
unchanged.

What lives here:
  - meta_get              — tolerant parsing of WooCommerce `meta_data`
  - WooClient             — authenticated, paginated REST client with a
                            retry-with-backoff layer and a truncation flag
  - detect_shop_currency  — read the shop currency from /system_status
  - iso_week_windows      — ISO-week-aligned current/prior reporting windows
  - upload_to_dropbox     — refresh-token OAuth upload (overwrite, muted)
  - Excel style helpers   — shared header styling / column widths

Nothing here is store-specific. Reports build on top of it.
"""

from __future__ import annotations

import json
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from typing import Any, Iterator

import requests
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Generic, vendor-agnostic helpers live in the dependency-light core package and
# are re-exported so consumers keep importing them from `wc_client` unchanged.
from code_commons import (  # noqa: F401  (re-exported for backward compatibility)
    CURRENCY_SYMBOLS,
    build_remote_path,
    currency_symbol,
    env_float,
    env_get,
    env_int,
    env_opt,
    env_required,
    log,
    parse_num,
)


# ---------------------------------------------------------------------------
# Tolerant WooCommerce meta parsing
# ---------------------------------------------------------------------------


def meta_get(meta_list: list | None, *keys: str) -> Any:
    """Return the first non-empty value for any of the given meta keys.

    Tolerates non-dict entries: some WC plugins write a bare string or other
    shape into `meta_data` alongside the normal `{"key", "value"}` dicts. A
    single malformed entry must not crash a run, so non-dicts are skipped.
    """
    if not meta_list:
        return None
    for key in keys:
        for m in meta_list:
            if not isinstance(m, dict):
                continue
            if m.get("key") == key:
                val = m.get("value")
                if val not in (None, "", []):
                    return val
    return None


# ---------------------------------------------------------------------------
# Currency detection (WooCommerce /system_status)
# ---------------------------------------------------------------------------


def detect_shop_currency(client: "WooClient") -> str:
    """Read the shop currency (3-letter code) from WooCommerce /system_status.

    Falls back to 'EUR' on any error — auto-detect failure must never block a
    run. Goes through the retry layer so a transient hiccup at startup does
    not regress to the fallback.
    """
    try:
        r = client.get_with_retry(f"{client.base}/system_status")
        currency = (r.json().get("settings", {}).get("currency") or "").upper()
        if not currency:
            log("WARNING: shop currency not in /system_status response; using EUR.")
            return "EUR"
        return currency
    except Exception as exc:  # noqa: BLE001 — fail-soft, never block the run
        log(f"WARNING: shop currency auto-detect failed ({exc}); using EUR.")
        return "EUR"


# ---------------------------------------------------------------------------
# WooCommerce client
# ---------------------------------------------------------------------------


class WooClient:
    """Authenticated, paginated WooCommerce REST v3 client.

    Parameter-based (not coupled to any report's Config) so it can be reused
    across routines.
    """

    def __init__(
        self,
        wc_url: str,
        consumer_key: str,
        consumer_secret: str,
        *,
        per_page: int = 100,
        max_pages: int = 50,
        timeout_seconds: int = 30,
        max_retries: int = 3,
        retry_backoff_base: float = 1.0,
    ):
        self.session = requests.Session()
        self.session.auth = (consumer_key, consumer_secret)
        self.base = f"{wc_url.rstrip('/')}/wp-json/wc/v3"
        self.per_page = per_page
        self.max_pages = max_pages
        self.timeout_seconds = timeout_seconds
        self.max_retries = max_retries
        self.retry_backoff_base = retry_backoff_base
        # F-005: set True when any paged() call stops at max_pages, so a caller
        # can tell a truncated result from a naturally-exhausted one and flag the
        # report as incomplete — a silent truncation would under-count silently.
        self.truncated = False

    def get_with_retry(self, url: str, params: dict | None = None):
        """GET with retry-with-backoff for transient errors.

        Retries on connection errors/timeouts, HTTP 5xx, and HTTP 429
        (rate-limited). Does NOT retry other 4xx — those are configuration
        errors (401 invalid key, 404 wrong endpoint) that retrying cannot fix.
        Backs off exponentially: retry_backoff_base * 2**attempt seconds.
        """
        last_exc: Exception | None = None
        for attempt in range(self.max_retries + 1):
            try:
                r = self.session.get(url, params=params, timeout=self.timeout_seconds)
                if r.status_code == 429 or 500 <= r.status_code < 600:
                    r.raise_for_status()  # raise so the retry path triggers
                r.raise_for_status()  # other 4xx falls through to caller, NOT retried
                return r
            except requests.HTTPError as exc:
                last_exc = exc
                code = exc.response.status_code if exc.response is not None else None
                if code is not None and code != 429 and not (500 <= code < 600):
                    raise  # non-retryable 4xx — bubble immediately
                if attempt >= self.max_retries:
                    raise
                wait = self.retry_backoff_base * (2 ** attempt)
                log(f"WARNING: WC {code} on {url} — retry {attempt + 1}/"
                    f"{self.max_retries} after {wait}s")
                time.sleep(wait)
            except (requests.ConnectionError, requests.Timeout) as exc:
                last_exc = exc
                if attempt >= self.max_retries:
                    raise
                wait = self.retry_backoff_base * (2 ** attempt)
                log(f"WARNING: WC connection error ({type(exc).__name__}) on "
                    f"{url} — retry {attempt + 1}/{self.max_retries} after {wait}s")
                time.sleep(wait)
        raise last_exc  # type: ignore[misc]  # unreachable: loop returns or raises

    def paged(self, path: str, params: dict | None = None) -> Iterator[dict]:
        """Yield every item across pages of a WC list endpoint, stopping after
        `max_pages` with a clear warning AND setting `self.truncated = True`, so a
        runaway loop or a much larger dataset is both visible and *detectable* by
        the caller rather than a silent under-count."""
        params = dict(params or {})
        params["per_page"] = self.per_page
        page = 1
        while True:
            if page > self.max_pages:
                self.truncated = True
                log(f"WARNING: reached MAX_PAGES={self.max_pages} on {path}. "
                    "Increase MAX_PAGES if your dataset is larger.")
                return
            params["page"] = page
            r = self.get_with_retry(f"{self.base}{path}", params=params)
            batch = r.json()
            if not batch:
                return
            for item in batch:
                yield item
            if len(batch) < self.per_page:
                return
            page += 1


# ---------------------------------------------------------------------------
# Dropbox upload
# ---------------------------------------------------------------------------


def upload_to_dropbox(
    *,
    app_key: str,
    app_secret: str,
    refresh_token: str,
    local_path: str,
    remote_path: str,
    timeout_seconds: int = 30,
) -> str:
    """Upload a local file to Dropbox via the app-key + refresh-token OAuth
    flow, in overwrite mode and muted (no desktop notification — these are
    automated daily runs). Returns the remote path on success; raises on
    failure so the caller can decide the exit code."""
    r = requests.post(
        "https://api.dropbox.com/oauth2/token",
        data={"grant_type": "refresh_token", "refresh_token": refresh_token},
        auth=(app_key, app_secret),
        timeout=timeout_seconds,
    )
    r.raise_for_status()
    access_token = r.json()["access_token"]

    with open(local_path, "rb") as f:
        data = f.read()
    args = {
        "path": remote_path,
        "mode": "overwrite",
        "autorename": False,
        "mute": True,
        "strict_conflict": False,
    }
    headers = {
        "Authorization": f"Bearer {access_token}",
        "Dropbox-API-Arg": json.dumps(args),
        "Content-Type": "application/octet-stream",
    }
    r = requests.post(
        "https://content.dropboxapi.com/2/files/upload",
        headers=headers,
        data=data,
        timeout=timeout_seconds * 2,
    )
    r.raise_for_status()
    return r.json().get("path_display", remote_path)


# ---------------------------------------------------------------------------
# ISO-week reporting windows
# ---------------------------------------------------------------------------
# Shared so reporting routines align "this period vs the same period
# last year" identically. ISO weeks always start on Monday, so aligning last
# year's window by ISO week number lines the two up by weekday automatically.


@dataclass
class Window:
    """A resolved reporting window. `before` is exclusive. `iso_start` is the
    ISO (year, week) of the window's first day."""
    after: datetime
    before: datetime
    label: str
    iso_start: tuple

    @property
    def days(self) -> int:
        return (self.before - self.after).days


def _utc_midnight(d: date) -> datetime:
    return datetime(d.year, d.month, d.day, tzinfo=timezone.utc)


def _monday_of_iso_week(d: date) -> date:
    """Return the Monday (ISO weekday 1) of the ISO week containing `d`."""
    iso_year, iso_week, _ = d.isocalendar()
    return date.fromisocalendar(iso_year, iso_week, 1)


def _safe_iso_monday(iso_year: int, iso_week: int) -> date:
    """Monday of (iso_year, iso_week), clamping week 53 to 52 when the target
    ISO year has no week 53 (no crash on the boundary)."""
    try:
        return date.fromisocalendar(iso_year, iso_week, 1)
    except ValueError:
        return date.fromisocalendar(iso_year, 52, 1)


def _window_label(start: date, end: date, iso_year: int, iso_week: int, weeks: int) -> str:
    last_day = end - timedelta(days=1)
    return (f"{iso_year}-W{iso_week:02d} +{weeks}wk "
            f"({start.isoformat()}–{last_day.isoformat()})")


def iso_week_windows(weeks: int, now: "datetime | None" = None) -> tuple[Window, Window]:
    """Resolve a current and a prior (one-ISO-year-earlier) window of `weeks`
    completed ISO weeks.

    Current: ends at 00:00 UTC on the Monday of the ISO week containing `now`
    (exclusive — the in-progress week is never half-counted) and starts `weeks`
    weeks earlier (also a Monday). Prior: the same ISO week number in the
    previous ISO year, weekday-aligned (both start on a Monday) and the same
    length. Week 53 clamps to 52 when last year has no week 53. `now` is
    injectable for tests.
    """
    if now is None:
        now = datetime.now(timezone.utc)

    current_end_date = _monday_of_iso_week(now.date())     # start of the in-progress week
    current_start_date = current_end_date - timedelta(weeks=weeks)
    cy, cw, _ = current_start_date.isocalendar()
    current = Window(
        after=_utc_midnight(current_start_date),
        before=_utc_midnight(current_end_date),
        label=_window_label(current_start_date, current_end_date, cy, cw, weeks),
        iso_start=(cy, cw),
    )

    prior_start_date = _safe_iso_monday(cy - 1, cw)
    prior_end_date = prior_start_date + timedelta(weeks=weeks)
    py, pw, _ = prior_start_date.isocalendar()
    prior = Window(
        after=_utc_midnight(prior_start_date),
        before=_utc_midnight(prior_end_date),
        label=_window_label(prior_start_date, prior_end_date, py, pw, weeks),
        iso_start=(py, pw),
    )
    return current, prior


# ---------------------------------------------------------------------------
# Goal / KPI arithmetic
# ---------------------------------------------------------------------------
# Pure, parameterised analytics shared by the reporting routines so the same
# goal numbers are computed identically everywhere: one definition of
# "year-over-year %", "house-brand revenue share" and "revenue vs target". They
# take plain numbers / {brand: revenue} maps and a house-brand *label* — no brand
# name or shop-specific meta key is baked in, so they stay vendor-agnostic. The
# caller owns rounding and presentation.


def yoy_pct(curr: float, prev: float) -> "float | None":
    """Year-over-year percentage: ``(curr - prev) / prev * 100``.

    Returns ``None`` when the prior is 0 or negative — growth is undefined for a
    genuinely new seller rather than an infinite spike. Naturally yields ``-100``
    when ``curr`` is 0 and ``prev`` > 0.
    """
    if prev <= 0:
        return None
    return (curr - prev) / prev * 100


def house_brand_share(
    curr_by_brand: "dict[str, float]",
    prev_by_brand: "dict[str, float]",
    house_brand: "str | None",
) -> dict:
    """House-brand revenue and its share of total revenue, for a current and a
    prior window, from two ``{brand: revenue}`` maps.

    The house brand is matched by a single case-insensitive label
    (``.strip().lower()``) — no brand name is hard-coded here; the caller passes
    the configured label. An empty/``None`` label yields zero house-brand revenue.
    Shares are ``None`` when the corresponding total is 0 (undefined, not zero).
    Values are returned unrounded; the caller rounds for display.
    """
    hb = (house_brand or "").strip().lower()
    total_curr = sum(curr_by_brand.values())
    total_prev = sum(prev_by_brand.values())
    hb_curr = sum(v for k, v in curr_by_brand.items() if k.strip().lower() == hb) if hb else 0.0
    hb_prev = sum(v for k, v in prev_by_brand.items() if k.strip().lower() == hb) if hb else 0.0
    return {
        "total_curr": total_curr,
        "total_prev": total_prev,
        "hb_curr": hb_curr,
        "hb_prev": hb_prev,
        "hb_share_curr": (hb_curr / total_curr * 100) if total_curr > 0 else None,
        "hb_share_prev": (hb_prev / total_prev * 100) if total_prev > 0 else None,
    }


def revenue_goal(
    total_curr: float,
    total_prev: float,
    *,
    target_growth_pct: "float | None" = None,
    target_absolute: "float | None" = None,
) -> dict:
    """Revenue versus target for one window.

    A growth-% target (``total_prev * (1 + pct/100)``) takes precedence over an
    absolute target; with neither, the target is undefined. A growth target needs
    a positive prior to resolve. ``revenue_target_pct`` is the share of target
    achieved (``curr / target * 100``). Values are unrounded; the caller rounds.
    """
    revenue_yoy = yoy_pct(total_curr, total_prev)
    if target_growth_pct is not None:
        target = total_prev * (1 + target_growth_pct / 100) if total_prev > 0 else None
        basis = f"+{target_growth_pct:g}% vs prior year"
    elif target_absolute is not None:
        target = target_absolute
        basis = "absolute"
    else:
        target, basis = None, None
    return {
        "revenue_yoy": revenue_yoy,
        "revenue_target": target,
        "revenue_target_basis": basis,
        "revenue_target_pct": (total_curr / target * 100) if target else None,
    }


# ---------------------------------------------------------------------------
# Excel style helpers
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
RED = PatternFill("solid", fgColor="F8CBAD")
ORANGE = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="C6E0B4")


def style_header(ws, n_cols: int) -> None:
    for col_idx in range(1, n_cols + 1):
        c = ws.cell(row=1, column=col_idx)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
